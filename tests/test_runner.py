"""Gesamtablauf mit nachgebildeter Modell-Laufzeit und echtem Tesseract/Pandoc (falls vorhanden)."""

import shutil
import socket
import urllib.request

import pytest
from docx import Document

from quellenuebersetzer import netguard
from quellenuebersetzer.catalog import load_catalog
from quellenuebersetzer.config import Paths, Settings
from quellenuebersetzer.pipeline.runner import Job, JobOptions
from quellenuebersetzer.profile import Profile, SourceKind
from quellenuebersetzer.runtime.ollama import OllamaClient

from conftest import make_text_page
from fake_ollama import FakeOllama


def test_netguard_blocks_outside_but_allows_localhost():
    with FakeOllama() as fake:
        with netguard.offline():
            assert OllamaClient(port=fake.port).is_up()
            with pytest.raises(netguard.NetworkBlocked):
                socket.create_connection(("192.0.2.1", 80), timeout=1)
            with pytest.raises(Exception):
                urllib.request.urlopen("http://example.org", timeout=2)
    assert not netguard.is_active()


def _job(root, fake, tmp_path, formats=("docx",), kind=SourceKind.DRUCK, events=None, layout="fliesstext"):
    settings = Settings(main_model="gemma-test", ocr_model="glm-test")
    profile = Profile([kind], "pl", ["18jh"])
    events = events if events is not None else []
    return Job(Paths(root), load_catalog(), settings, profile,
               JobOptions(kind, list(formats), review=False, output_root=tmp_path / "out", layout=layout),
               emit=events.append, client=OllamaClient(port=fake.port))


def test_full_job_on_synthetic_scan(tmp_path, tess_root):
    src = make_text_page(tmp_path / "Akte 1785.png")
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, events=events).run([src])
    assert results[0].ok, results[0].message
    names = sorted(p.name for p in results[0].outputs)
    assert "Akte 1785 – Übersetzung.docx" in names and "Akte 1785 – Umschrift.md" in names
    de_md = next(p for p in results[0].outputs if p.name.endswith("Übersetzung.md")).read_text(encoding="utf-8")
    assert "Krakau" in de_md and "Warnhinweis Handschrift" not in de_md
    assert "Begriff der Zeit" in de_md  # B2 aus der (nachgebildeten) Übersetzung im Verzeichnis
    doc = Document(str(next(p for p in results[0].outputs if p.name.endswith("Übersetzung.docx"))))
    assert len(doc.tables) == 1 and "Bronowice" in doc.tables[0].cell(1, 0).text
    assert "translate" in fake.calls
    # Seitenbilder wurden nach der Verarbeitung entfernt
    assert not any((tess_root / "cache" / "jobs").glob("*/*.png"))
    assert any(e.kind == "fertig" for e in events)


def test_handwriting_document_carries_warning(tmp_path, tess_root):
    src = make_text_page(tmp_path / "brief.png", with_table=False)
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=(), kind=SourceKind.HANDSCHRIFT).run([src])
    assert results[0].ok, results[0].message
    md = next(p for p in results[0].outputs if p.name.endswith("Übersetzung.md")).read_text(encoding="utf-8")
    assert "Warnhinweis Handschrift" in md and "chronisch unzuverlässig" in md
    assert "htr" in fake.calls


def test_digital_docx_is_translated_without_ocr(tmp_path, tess_root):
    d = Document()
    d.add_heading("Przywilej", level=1)
    d.add_paragraph("Starosta grodowy zatwierdził przywilej dla cechu młynarzy.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Miejscowość", "Młyny"
    t.cell(1, 0).text, t.cell(1, 1).text = "Bronowice", "3"
    src = tmp_path / "akte.docx"
    d.save(str(src))
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=()).run([src])
    assert results[0].ok
    assert "ocr" not in fake.calls
    md = next(p for p in results[0].outputs if p.name.endswith("Übersetzung.md")).read_text(encoding="utf-8")
    assert "Mühlen" in md and "<table>" in md


def test_helper_processes_get_no_outside_network(tmp_path):
    """Hilfsprozesse (Ollama, Pandoc, Typst, Tesseract) laufen mit Umgebung, die nur 127.0.0.1 erreicht."""
    import subprocess
    import sys

    from quellenuebersetzer.netguard import offline_env
    from quellenuebersetzer.runtime.ollama import OllamaService

    probe = (
        "import sys, urllib.request\n"
        "try:\n"
        "    urllib.request.urlopen('http://example.org', timeout=5); print('DRAUSSEN')\n"
        "except Exception:\n"
        "    print('GESPERRT')\n"
        "urllib.request.urlopen(sys.argv[1], timeout=5); print('LOKAL')\n"
    )
    with FakeOllama() as fake:
        out = subprocess.run([sys.executable, "-c", probe, f"http://127.0.0.1:{fake.port}/api/version"],
                             env=offline_env(), capture_output=True, text=True, timeout=60).stdout
    assert "GESPERRT" in out and "LOKAL" in out and "DRAUSSEN" not in out

    env = OllamaService(Paths(tmp_path)).environment(offline=True)
    assert env["HTTPS_PROXY"].startswith("http://127.0.0.1:") and env["OLLAMA_NO_CLOUD"] == "1"
    assert "127.0.0.1" in env["NO_PROXY"]
    assert "HTTPS_PROXY" not in OllamaService(Paths(tmp_path)).environment(offline=False) or \
        not OllamaService(Paths(tmp_path)).environment(offline=False)["HTTPS_PROXY"].startswith("http://127.0.0.1:9")


def test_docx_source_is_translated_in_place_when_page_faithful(tmp_path, tess_root):
    d = Document()
    d.add_heading("Przywilej", level=1)
    d.add_paragraph("Starosta grodowy zatwierdził przywilej dla cechu młynarzy.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Miejscowość", "Młyny"
    t.cell(1, 0).text, t.cell(1, 1).text = "Bronowice", "3"
    src = tmp_path / "akte.docx"
    d.save(str(src))
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("docx",), layout="faksimile").run([src])
    assert results[0].ok, results[0].message
    out = Document(str(results[0].outputs[0]))
    assert out.paragraphs[0].style.name.startswith("Heading")  # Formatvorlage erhalten
    assert "bestätigte" in out.paragraphs[1].text and "Starost" in out.paragraphs[1].text
    assert out.tables[0].cell(0, 0).text == "Ort" and out.tables[0].cell(1, 1).text == "3"
    assert any("Anmerkungen" in p.text for p in out.paragraphs)  # B2-Erläuterung im Anhang


def test_layout_job_on_mixed_page(tmp_path, tess_root):
    from conftest import make_mixed_page
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    src = make_mixed_page(tmp_path / "Chronik.png")
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("pdf", "docx"), layout="faksimile", events=events).run([src])
    assert results[0].ok, results[0].message
    previews = [e for e in events if e.kind == "seite_fertig"]
    assert len(previews) == 1 and all(p.exists() for p in previews[0].outputs)  # Vorschau der fertigen Seite
    names = {p.name for p in results[0].outputs}
    assert {"Chronik – Übersetzung.pdf", "Chronik – Übersetzung.docx", "Chronik – Umschrift.pdf"} <= names
    md = next(p for p in results[0].outputs if p.name == "Chronik – Übersetzung.md").read_text(encoding="utf-8")
    assert "Krakau" in md and "Abbildungen/abbildung-" in md
    assert (tmp_path / "out" / "Chronik" / "Abbildungen").is_dir()
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(tmp_path / "out" / "Chronik" / "Chronik – Übersetzung.pdf"))
    assert len(pdf) == 2  # Seite + Anhang
    w, h = pdf[0].get_size()
    assert abs(w - 2480 / 300 * 72) < 2 and abs(h - 3508 / 300 * 72) < 2  # Originalformat


def test_layout_job_resumes_after_interruption(tmp_path, tess_root):
    from conftest import make_mixed_page
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    src = make_mixed_page(tmp_path / "Band.png")
    events = []
    with FakeOllama() as fake:
        job = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events)

        def emit(ev, _orig=job.emit):
            _orig(ev)
            if ev.kind == "schritt" and ev.message.startswith("Phase 3 von 3 – Übersetzung: Abschnitt 2 "):
                job.cancel()  # Unterbrechung mitten in der Übersetzung

        job.emit = emit
        first = job.run([src])
        assert not first[0].ok
        assert list((tess_root / "cache" / "buecher").glob("*/stand.json"))  # Zwischenstand erhalten

        events2 = []
        second = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events2).run([src])
    assert second[0].ok, second[0].message
    assert any("fortgesetzt" in e.message for e in events2 if e.kind == "hinweis")
    assert not any(e.message.startswith(("Phase 1", "Phase 2")) for e in events2 if e.kind == "schritt")
    book_dirs = [d for d in (tess_root / "cache" / "buecher").iterdir()] if (tess_root / "cache" / "buecher").exists() else []
    assert not any((d / "stand.json").exists() for d in book_dirs)  # nach Abschluss aufgeräumt


def test_scans_use_layout_first_also_for_flow_text(tmp_path, tess_root):
    import zipfile

    from conftest import make_mixed_page
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    src = make_mixed_page(tmp_path / "Chronik.png")
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("docx",), layout="fliesstext", events=events).run([src])
    assert results[0].ok, results[0].message
    assert any(e.message.startswith("Phase 1 von 3 – Seitenaufbau") for e in events if e.kind == "schritt")
    docx_path = next(p for p in results[0].outputs if p.name == "Chronik – Übersetzung.docx")
    with zipfile.ZipFile(docx_path) as z:
        assert any(n.startswith("word/media/") for n in z.namelist())  # Abbildung übernommen
        xml = z.read("word/document.xml").decode("utf-8")
    assert "Randnotiz" in xml and "framePr" not in xml  # Fließtext, keine positionierten Rahmen


def test_preview_shows_read_pages_first_and_timing_log_has_no_names(tmp_path, tess_root):
    from conftest import make_mixed_page
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    src = make_mixed_page(tmp_path / "Geheime Akte Kowalski.png")
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events).run([src])
    assert results[0].ok, results[0].message
    kinds = [e.kind for e in events if e.kind in ("seite_gelesen", "seite_fertig")]
    assert kinds == ["seite_gelesen", "seite_fertig"]  # dann gelesen, dann übersetzt
    first = next(i for i, e in enumerate(events) if e.kind == "seite_aufbau")
    assert first < next(i for i, e in enumerate(events) if e.kind == "seite_gelesen")  # Aufbau zuerst
    log = (tess_root / "logs" / "zeiten.log").read_text(encoding="utf-8")
    last = log.strip().split(" | Quelle ")[-1]
    assert "Kowalski" not in log and "Geheime" not in log
    assert "Lesen:" in last and "Übersetzen:" in last and "gemma-test" in last


def test_preview_panel_replaces_read_page_with_translation(tmp_path):
    from PySide6.QtWidgets import QApplication

    from quellenuebersetzer.ui.preview import PreviewPanel

    app = QApplication.instance() or QApplication(["test", "-platform", "offscreen"])
    panel = PreviewPanel()
    img = tmp_path / "a.png"
    from PIL import Image

    Image.new("RGB", (40, 60), "white").save(img)
    panel.add_page("1", 1, 0, img, None)
    panel.add_page("2", 2, 0, img, None)
    assert "wird übersetzt" in panel.stage_label.text()
    panel.add_page("1", 1, 2, img, img)
    assert len(panel.pages) == 2 and panel.pages[0].translated == img
    assert "übersetzt 1" in panel.position.text()


def test_whole_document_layout_comes_before_any_reading(tmp_path, tess_root):
    from conftest import make_mixed_page, make_text_page
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    book = tmp_path / "Band"
    book.mkdir()
    make_mixed_page(book / "s1.png")
    make_text_page(book / "s2.png", with_table=False)
    make_text_page(book / "s3.png")
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events).run([book])
    assert results[0].ok, results[0].message
    seq = [e.kind for e in events if e.kind.startswith("seite_")]
    last_layout = max(i for i, k in enumerate(seq) if k == "seite_aufbau")
    first_read = min(i for i, k in enumerate(seq) if k == "seite_gelesen")
    assert seq.count("seite_aufbau") == 3 and last_layout < first_read  # erst alle Seiten Aufbau, dann Lesen
    assert any(e.kind == "hinweis" and e.message.startswith("Seitenaufbau erkannt: 3 Seiten") for e in events)


def test_distorted_scans_are_straightened_before_layout(tmp_path, tess_root):
    from conftest import make_text_page
    from quellenuebersetzer.runtime.layout import LayoutDetector
    from test_tables_reading import make_distorted_scan

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    book = tmp_path / "Akte"
    book.mkdir()
    make_distorted_scan(book / "s1.png")
    make_text_page(book / "s2.png", with_table=False, size=(1650, 2350))
    events = []
    with FakeOllama() as fake:
        results = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events).run([book])
    assert results[0].ok, results[0].message
    assert any(e.kind == "hinweis" and "entzerrt" in e.message for e in events)
    md = next(p for p in results[0].outputs if p.name.endswith("Umschrift.md")).read_text(encoding="utf-8")
    assert md.count("Starosta") >= 5  # die entzerrte Seite wird lesbar erkannt
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(results[0].outputs[-1]))
    sizes = {tuple(round(v) for v in pdf[i].get_size()) for i in range(2)}
    assert len(sizes) == 1  # einheitliches Seitenformat


def test_saved_state_of_another_program_version_is_not_reused(tmp_path, tess_root):
    """Nach einem Update wird der Seitenaufbau neu erkannt – kein alter Zwischenstand."""
    from conftest import make_mixed_page
    from quellenuebersetzer.pipeline.layout_doc import LayoutDocument
    from quellenuebersetzer.runtime.layout import LayoutDetector

    if not LayoutDetector(Paths(tess_root)).available:
        pytest.skip("Seitenaufbau-Modell nicht in der Testumgebung")
    src = make_mixed_page(tmp_path / "Band.png")
    events = []
    with FakeOllama() as fake:
        job = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events)

        def emit(ev, _orig=job.emit):
            _orig(ev)
            if ev.kind == "schritt" and ev.message.startswith("Phase 3 von 3 – Übersetzung: Abschnitt 1 "):
                job.cancel()

        job.emit = emit
        job.run([src])
        state = next(job.paths.cache.rglob("stand.json"))
        doc = LayoutDocument.load(state)
        doc.programm = "0.0.1 (Bau 01.01.2026 00:00)"  # Zwischenstand einer älteren Fassung
        doc.save(state)
        events2 = []
        second = _job(tess_root, fake, tmp_path, formats=("pdf",), layout="neusatz", events=events2).run([src])
    assert second[0].ok
    assert any("aktualisiert" in e.message for e in events2 if e.kind == "hinweis")
    assert not any("fortgesetzt" in e.message for e in events2 if e.kind == "hinweis")
