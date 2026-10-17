"""Seitengetreue Word-Datei: jede Seite ein Abschnitt in Originalgröße, jeder Bereich ein positionierter Rahmen.

- Faksimile: retuschierter Scan als Bild hinter dem Text; Neusatz: Abbildungen als verankerte Bilder.
- Text: Absatzrahmen (``w:framePr``) mit exakter Position und Höhe; die Schriftgröße wird vorab mit den
  echten Schriftmaßen so bestimmt, dass der Text in den Rahmen passt (Word kann das nicht selbst).
- Tabellen: schwebende Tabellen (``w:tblpPr``) mit exakten Zeilenhöhen.
- Markierungen: Zeichenformatvorlagen wie in der Fließtextfassung + hochgestellte Anmerkungsnummer;
  Erläuterungen gesammelt im Anhang „Anmerkungen“.
"""

from __future__ import annotations

import html
from dataclasses import dataclass
import os
import re
from functools import lru_cache
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt
from PIL import ImageFont

from quellenuebersetzer.config import resources_dir
from quellenuebersetzer.pipeline import markup
from quellenuebersetzer.pipeline.layout_doc import HEADING_KINDS, IMAGE_KINDS, LayoutDocument, LPage, LRegion
from quellenuebersetzer.pipeline.fitting import expand_boxes, fit_text
from quellenuebersetzer.pipeline.layout_render import MIN_SCALE, Notes, erase_text
from quellenuebersetzer.planner import GENERAL_DISCLAIMER, HANDWRITING_WARNING

FONT = "Times New Roman"
STYLE_NAMES = {"a": "Lesung unsicher", "b1": "Bedeutungswandel", "b2": "Begriff der Zeit"}
EMU_PER_PT = 12700
TWIPS_PER_PT = 20
LINE_FACTOR = 1.15


def _px_to_pt(px: float, dpi: float) -> float:
    return px / dpi * 72.0


# --------------------------------------------------------------------------- Schriftmaße


@lru_cache(maxsize=8)
def _font(bold: bool, size_tenths: int) -> ImageFont.FreeTypeFont:
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    for name in (("timesbd.ttf", "georgiab.ttf") if bold else ("times.ttf", "georgia.ttf")):
        if (fonts / name).exists():
            return ImageFont.truetype(str(fonts / name), max(1, size_tenths) / 10)
    return ImageFont.load_default(max(1, size_tenths) / 10)


def text_height(text: str, width_pt: float, size_pt: float, bold: bool) -> float:
    """Höhe des umbrochenen Textes in pt (gieriger Zeilenumbruch, wie ihn Word annähernd macht)."""
    font = _font(bold, int(size_pt * 10))
    lines, current = 1, 0.0
    space = font.getlength(" ")
    for word in text.split():
        wl = font.getlength(word)
        if current and current + space + wl > width_pt * 0.97:
            lines += 1
            current = wl
        else:
            current += (space if current else 0) + wl
    return lines * size_pt * LINE_FACTOR


def fit_size(text: str, width_pt: float, height_pt: float, size_pt: float, bold: bool) -> float:
    s = size_pt
    while s > size_pt * MIN_SCALE and text_height(text, width_pt, s, bold) > height_pt:
        s *= 0.95
    return max(4.0, s)


# --------------------------------------------------------------------------- XML-Bausteine

@dataclass
class RunStyle:
    bold: bool = False
    italic: bool = False
    color: str = ""  # Hex
    tracking_twips: int = 0  # Laufweite in Zwanzigstel-Punkt


def _run(text: str, size_pt: float, bold: bool = False, style_id: str | None = None, superscript: bool = False,
         st: "RunStyle | None" = None) -> str:
    rpr = [f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}"/>']
    if style_id:
        rpr.insert(0, f'<w:rStyle w:val="{style_id}"/>')
    if bold or (st and st.bold):
        rpr.append("<w:b/>")
    if st and st.italic:
        rpr.append("<w:i/>")
    if st and st.color and not style_id:
        rpr.append(f'<w:color w:val="{st.color.lstrip("#").upper()}"/>')
    if st and st.tracking_twips:
        rpr.append(f'<w:spacing w:val="{st.tracking_twips}"/>')
    if superscript:
        rpr.append('<w:vertAlign w:val="superscript"/>')
    rpr.append(f'<w:sz w:val="{max(2, round(size_pt * 2))}"/><w:szCs w:val="{max(2, round(size_pt * 2))}"/>')
    return f'<w:r><w:rPr>{"".join(rpr)}</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def _runs_from_md(md: str, size: float, bold: bool, notes: Notes, style_ids: dict[str, str],
                  st: RunStyle | None = None) -> tuple[str, str]:
    """→ (Runs-XML, Klartext)."""
    out, plain, pos = [], [], 0
    for start, end, mark in markup.find_marks(md):
        t = markup.unescape_md(md[pos:start])
        if t:
            out.append(_run(t, size, bold, st=st))
            plain.append(t)
        out.append(_run(mark.text, size, bold, style_ids.get(mark.category), st=st))
        out.append(_run(str(notes.add(mark)), size, bold, superscript=True, st=st))
        plain.append(mark.text + "¹")
        pos = end
    t = markup.unescape_md(md[pos:])
    if t:
        out.append(_run(t, size, bold, st=st))
        plain.append(t)
    return "".join(out), "".join(plain)


def _frame_paragraph(runs: str, x: float, y: float, w: float, h: float, size: float, align: str,
                     leading: float = LINE_FACTOR) -> str:
    jc = {"left": "left", "center": "center", "right": "right", "justify": "both"}.get(align, "left")
    pad = 0.35 * size  # Luft oben für Akzente über Versalien (Ü, Ś, Ż)
    y, h = y - pad, h + pad
    return (
        f'<w:p {nsdecls("w")}><w:pPr>'
        f'<w:framePr w:w="{round(w * TWIPS_PER_PT)}" w:h="{round(h * TWIPS_PER_PT)}" w:hRule="exact" w:wrap="around" '
        f'w:hAnchor="page" w:vAnchor="page" w:x="{round(x * TWIPS_PER_PT)}" w:y="{round(y * TWIPS_PER_PT)}"/>'
        f'<w:spacing w:before="{round(pad * TWIPS_PER_PT)}" w:after="0" w:line="{round(size * leading * TWIPS_PER_PT)}" w:lineRule="exact"/>'
        f'<w:ind w:left="0" w:right="0"/><w:jc w:val="{jc}"/></w:pPr>{runs}</w:p>'
    )


def _anchor_image(r_id: str, x: float, y: float, w: float, h: float, pic_id: int, behind: bool) -> str:
    cx, cy, ox, oy = (round(v * EMU_PER_PT) for v in (w, h, x, y))
    return (
        f'<w:r {nsdecls("w", "wp", "a", "pic", "r")}><w:drawing>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{pic_id}" '
        f'behindDoc="{1 if behind else 0}" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{ox}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{oy}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{pic_id}" name="Bild {pic_id}"/><wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{pic_id}" name="bild{pic_id}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{r_id}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
        f'</wp:anchor></w:drawing></w:r>'
    )


def _table_xml(table_html: str, x: float, y: float, w: float, h: float, notes: Notes, style_ids: dict[str, str]) -> str:
    rows = re.findall(r"<tr\b[^>]*>(.*?)</tr>", table_html, re.S | re.I)
    grid: list[list[tuple[str, int, int, bool]]] = []
    n_cols = 0
    occupied: set[tuple[int, int]] = set()
    placed: dict[tuple[int, int], tuple[str, int, int, bool]] = {}
    for r, row in enumerate(rows):
        c = 0
        for m in re.finditer(r"<(t[dh])\b([^>]*)>(.*?)</t[dh]>", row, re.S | re.I):
            while (r, c) in occupied:
                c += 1
            rs = int((re.search(r'rowspan="(\d+)"', m.group(2)) or [0, 1])[1])
            cs = int((re.search(r'colspan="(\d+)"', m.group(2)) or [0, 1])[1])
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
            placed[(r, c)] = (m.group(3), rs, cs, m.group(1).lower() == "th")
            c += cs
            n_cols = max(n_cols, c)
    n_rows = len(rows)
    if not n_rows or not n_cols:
        return ""
    col_w = w / n_cols
    row_h = h / n_rows
    # gemeinsame Schriftgröße, die in alle Zellen passt
    size = min(11.0, row_h / LINE_FACTOR * 0.85)
    for (r, c), (inner, rs, cs, th) in placed.items():
        text = html.unescape(re.sub(r"<[^>]+>", "", inner))
        size = min(size, fit_size(text, col_w * cs - 8, row_h * rs - 2, size, th))
    xml_rows = []
    for r in range(n_rows):
        cells = []
        c = 0
        while c < n_cols:
            if (r, c) in placed:
                inner, rs, cs, th = placed[(r, c)]
                runs = _cell_runs(inner, size, th, notes, style_ids)
                vmerge = '<w:vMerge w:val="restart"/>' if rs > 1 else ""
                cells.append(f'<w:tc><w:tcPr><w:tcW w:w="{round(col_w * cs * TWIPS_PER_PT)}" w:type="dxa"/>'
                             f'{"<w:gridSpan w:val=" + chr(34) + str(cs) + chr(34) + "/>" if cs > 1 else ""}{vmerge}</w:tcPr>'
                             f'<w:p><w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr>{runs}</w:p></w:tc>')
                c += cs
            else:
                # Teil einer senkrecht verbundenen Zelle
                owner = next(((rr, cc) for (rr, cc), v in placed.items() if rr < r <= rr + v[1] - 1 and cc == c), None)
                cs = placed[owner][2] if owner else 1
                cells.append(f'<w:tc><w:tcPr><w:tcW w:w="{round(col_w * cs * TWIPS_PER_PT)}" w:type="dxa"/>'
                             f'{"<w:gridSpan w:val=" + chr(34) + str(cs) + chr(34) + "/>" if cs > 1 else ""}'
                             f'<w:vMerge/></w:tcPr><w:p/></w:tc>')
                c += cs
        xml_rows.append(f'<w:tr><w:trPr><w:trHeight w:val="{round(row_h * TWIPS_PER_PT)}" w:hRule="exact"/></w:trPr>'
                        + "".join(cells) + "</w:tr>")
    border = '<w:{0} w:val="single" w:sz="4" w:space="0" w:color="6B6255"/>'
    borders = "".join(border.format(s) for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
    grid_cols = "".join(f'<w:gridCol w:w="{round(col_w * TWIPS_PER_PT)}"/>' for _ in range(n_cols))
    return (
        f'<w:tbl {nsdecls("w")}><w:tblPr>'
        f'<w:tblpPr w:leftFromText="0" w:rightFromText="0" w:vertAnchor="page" w:horzAnchor="page" '
        f'w:tblpX="{round(x * TWIPS_PER_PT)}" w:tblpY="{round(y * TWIPS_PER_PT)}"/>'
        f'<w:tblW w:w="{round(w * TWIPS_PER_PT)}" w:type="dxa"/><w:tblLayout w:type="fixed"/>'
        f'<w:tblBorders>{borders}</w:tblBorders>'
        f'<w:tblCellMar><w:left w:w="57" w:type="dxa"/><w:right w:w="57" w:type="dxa"/></w:tblCellMar>'
        f'</w:tblPr><w:tblGrid>{grid_cols}</w:tblGrid>{"".join(xml_rows)}</w:tbl>'
    )


def _cell_runs(inner: str, size: float, bold: bool, notes: Notes, style_ids: dict[str, str]) -> str:
    out, pos = [], 0
    for m in markup.HTML_SPAN_RE.finditer(inner):
        t = html.unescape(re.sub(r"<[^>]+>", "", inner[pos:m.start()]))
        if t:
            out.append(_run(t, size, bold))
        attrs = {k: html.unescape(v) for k, v in re.findall(r'([\w-]+)="([^"]*)"', m.group("attrs") or "")}
        text = html.unescape(re.sub(r"<[^>]+>", "", m.group("text")))
        mark = markup.Mark(m.group("cat"), text, attrs)
        out.append(_run(text, size, bold, style_ids.get(mark.category)))
        out.append(_run(str(notes.add(mark)), size, bold, superscript=True))
        pos = m.end()
    t = html.unescape(re.sub(r"<[^>]+>", "", inner[pos:]))
    if t:
        out.append(_run(t, size, bold))
    return "".join(out)


# --------------------------------------------------------------------------- Dokument


def build_docx(doc: LayoutDocument, out: Path, faksimile: bool = True, translated: bool = True,
               handwriting: bool = False, title: str = "") -> Path:
    d = Document(str(resources_dir() / "export" / "referenz.docx"))
    body = d.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    style_ids = {}
    for cat, name in STYLE_NAMES.items():
        try:
            style_ids[cat] = d.styles[name].style_id
        except KeyError:
            pass
    notes = Notes()
    work = Path(doc.workdir)
    pic_id = 1000

    for idx, page in enumerate(doc.pages):
        notes.page = page.name
        section = d.sections[0] if idx == 0 else d.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Emu(round(_px_to_pt(page.width_px, page.dpi) * EMU_PER_PT))
        section.page_height = Emu(round(_px_to_pt(page.height_px, page.dpi) * EMU_PER_PT))
        for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance"):
            setattr(section, attr, Emu(0))
        anchor_p = _section_last_paragraph(d, idx)
        if faksimile:
            name = f"faksimile-{page.number:04d}-de.jpg" if translated else page.background
            if translated and not (work / name).exists():
                erase_text(work / page.background, page, work / name)
            r_id, _ = d.part.get_or_add_image(str(work / name))
            pic_id += 1
            anchor_p.append(parse_xml(_anchor_image(r_id, 0, 0, _px_to_pt(page.width_px, page.dpi),
                                                    _px_to_pt(page.height_px, page.dpi), pic_id, behind=True)))
        if faksimile and not translated:
            continue
        if not faksimile:
            for dec in page.decorations:
                png = panel_png(dec, page, work)
                r_id, _ = d.part.get_or_add_image(str(png))
                pic_id += 1
                x0, y0, x1, y1 = dec["box"]
                anchor_p.append(parse_xml(_anchor_image(
                    r_id, _px_to_pt(x0, page.dpi), _px_to_pt(y0, page.dpi), _px_to_pt(x1 - x0, page.dpi),
                    _px_to_pt(y1 - y0, page.dpi), pic_id, behind=True)))
        boxes = expand_boxes(page.regions, page.height_px, decorations=page.decorations)
        for r in page.regions:
            pic_id = _add_region(d, anchor_p, r, page, notes, style_ids, faksimile, translated, pic_id, work,
                                 boxes.get(r.id))

    # Anhang
    last = d.add_section(WD_SECTION.NEW_PAGE)
    last.page_width, last.page_height = Emu(595 * EMU_PER_PT), Emu(842 * EMU_PER_PT)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(last, attr, Pt(57))
    d.add_heading(title or doc.title, level=1)
    if handwriting:
        p = d.add_paragraph(style="Warnung") if "Warnung" in [s.name for s in d.styles] else d.add_paragraph()
        p.add_run("Warnhinweis Handschrift: ").bold = True
        p.add_run(HANDWRITING_WARNING)
    p = d.add_paragraph(style="Hinweis") if "Hinweis" in [s.name for s in d.styles] else d.add_paragraph()
    p.add_run(GENERAL_DISCLAIMER)
    if translated:
        from quellenuebersetzer.pipeline import pruefbericht

        rows = [(n, l, i) for n, l, i in pruefbericht.summary(doc) if i]
        if rows:
            d.add_heading("Prüfbericht", level=2)
            d.add_paragraph("Seiten mit Auffälligkeiten (rot = Fehler, gelb = prüfen). Alle übrigen Seiten ohne "
                            "Hinweise.")
            for name, l, issues in rows:
                p = d.add_paragraph()
                p.add_run(f"S. {name} ({l})").bold = True
                for issue in issues:
                    d.add_paragraph(f"– {issue}")
    if notes.shortened:
        d.add_heading("Vollständiger Text gekürzter Bereiche", level=2)
        d.add_paragraph("Diese Bereiche passten auch bei engem Satz nicht vollständig auf ihre Seite und sind dort "
                        "mit […] gekürzt.")
        for page_label, md in notes.shortened:
            p = d.add_paragraph()
            p.add_run(f"S. {page_label}: ").bold = True
            p.add_run(markup.unescape_md(markup.strip_marks(md)))
    if notes.items:
        d.add_heading("Anmerkungen", level=2)
        d.add_paragraph("Orange gepunktet = Lesung unsicher · blau = Bedeutungswandel · grün gestrichelt = Begriff der Zeit.")
        for page_label, n, text in notes.items:
            d.add_paragraph(f"S. {page_label}, Anm. {n}: {text}")
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def panel_png(dec: dict, page: LPage, work: Path) -> Path:
    """Kasten (Hintergrundfarbe, Rahmen) als Bild – liegt im Word-Neusatz hinter den Textrahmen."""
    from PIL import Image, ImageDraw

    out = work / f"kasten-{page.number:04d}-{dec['id']}.png"
    if out.exists():
        return out
    x0, y0, x1, y1 = dec["box"]
    scale = 0.5
    w, h = max(4, int((x1 - x0) * scale)), max(4, int((y1 - y0) * scale))
    img = Image.new("RGBA", (w, h), dec.get("fill") or (0, 0, 0, 0))
    if dec.get("stroke"):
        dr = ImageDraw.Draw(img)
        t = max(1, int((dec.get("stroke_px") or 3) * scale))
        dash, gap = {"dashed": (8, 5), "dotted": (2, 4)}.get(dec.get("dash") or "", (0, 0))
        for (ax, ay, bx, by) in ((0, 0, w, 0), (0, h - t, w, h - t), (0, 0, 0, h), (w - t, 0, w - t, h)):
            horizontal = ay == by
            length = (bx - ax) if horizontal else (by - ay)
            pos = 0
            while pos < length:
                end = length if not dash else min(length, pos + dash)
                if horizontal:
                    dr.rectangle([ax + pos, ay, ax + end, ay + t - 1], fill=dec["stroke"])
                else:
                    dr.rectangle([ax, ay + pos, ax + t - 1, ay + end], fill=dec["stroke"])
                pos = end + gap if dash else length
    img.save(out)
    return out


def _section_last_paragraph(d, idx: int):
    """Absatz, der den Abschnitt der aktuellen Seite trägt (Anker für Bilder und Rahmen)."""
    body = d.element.body
    if idx == 0:
        p = parse_xml(f'<w:p {nsdecls("w")}/>')
        body.insert(0, p)
        return p
    # add_section hat einen Absatz mit Abschnittsumbruch angehängt; der neue Inhalt kommt danach
    p = parse_xml(f'<w:p {nsdecls("w")}/>')
    body.insert(len(body) - 1, p)
    return p


def _add_region(d, anchor_p, r: LRegion, page: LPage, notes: Notes, style_ids, faksimile: bool, translated: bool,
                pic_id: int, work: Path, box_px=None) -> int:
    dpi = page.dpi
    x0, y0, x1, y1 = box_px or r.box
    x, y, w, h = _px_to_pt(x0, dpi), _px_to_pt(y0, dpi), _px_to_pt(x1 - x0, dpi), _px_to_pt(y1 - y0, dpi)
    if r.kind in IMAGE_KINDS:
        if not faksimile and r.image:
            r_id, _ = d.part.get_or_add_image(str(work / r.image))
            pic_id += 1
            anchor_p.append(parse_xml(_anchor_image(r_id, x, y, w, h, pic_id, behind=False)))
        return pic_id
    text = r.target_md if translated else r.source_md
    if not text.strip():
        return pic_id
    if r.kind == "tabelle":
        anchor_p.addprevious(parse_xml(_table_xml(text, x, y, w, h, notes, style_ids)))
        anchor_p.addprevious(parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="0" w:after="0"/>'
                                       f'<w:rPr><w:sz w:val="2"/></w:rPr></w:pPr></w:p>'))
        return pic_id
    bold = r.bold or r.kind in HEADING_KINDS
    base = max(5.0, min(40.0, _px_to_pt(r.line_px or (y1 - y0), dpi) * 0.92))
    fit = fit_text(text, (x, y, w, h), base, bold, r.italic, r.spaced)
    if fit.truncated:
        notes.shortened.append((page.name, text))
    st = RunStyle(bold, r.italic, r.color, round(fit.tracking * fit.size * TWIPS_PER_PT))
    runs, _ = _runs_from_md(fit.text_md, fit.size, bold, notes, style_ids, st)
    anchor_p.addprevious(parse_xml(_frame_paragraph(runs, x, y, w, h, fit.size, r.align, fit.leading)))
    return pic_id


render_docx = build_docx
